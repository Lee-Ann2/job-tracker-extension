import pandas as pd
from datetime import datetime
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

class ExcelManager:
    def __init__(self, filename="job_applications.xlsx"):
        self.filename = filename
        self.columns = [
            "Date Applied", "Company", "Position", "Platform", 
            "Status", "Application URL", "Notes", "Last Updated"
        ]
        self.status_colors = {
            "Applied": "FFE4B5",  
            "Interview": "C6EFCE",
            "Rejected": "FFC7CE",
            "Ignored": "FFEB9C",
            "Offer": "C6EFCE"  
        }
        self._initialize_excel()
    
    def _initialize_excel(self):
        """Create Excel file if it doesn't exist"""
        if not os.path.exists(self.filename):
            df = pd.DataFrame(columns=self.columns)
            df.to_excel(self.filename, index=False)
            self._apply_styling()
    
    def _apply_styling(self):
        """Apply styling to Excel file"""
        wb = load_workbook(self.filename)
        ws = wb.active
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        
        ws.column_dimensions['A'].width = 15  
        ws.column_dimensions['B'].width = 25  
        ws.column_dimensions['C'].width = 30  
        ws.column_dimensions['D'].width = 12  
        ws.column_dimensions['E'].width = 12  
        ws.column_dimensions['F'].width = 50  
        ws.column_dimensions['G'].width = 40  
        ws.column_dimensions['H'].width = 15  
        
        wb.save(self.filename)
    
    def add_application(self, company, position, platform, url, notes=""):
        """Add a new job application"""
        df = pd.read_excel(self.filename)
        
        new_entry = {
            "Date Applied": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "Company": company,
            "Position": position,
            "Platform": platform,
            "Status": "Applied",
            "Application URL": url,
            "Notes": notes,
            "Last Updated": datetime.now().strftime("%Y-%m-%d %H:%M")
        }
        
        df = pd.concat([df, pd.DataFrame([new_entry])], ignore_index=True)
        df.to_excel(self.filename, index=False)
        self._update_status_color(len(df) - 1)
        
        return len(df) - 1  
    
    def update_status(self, row_index, status, notes=""):
        """Update application status"""
        df = pd.read_excel(self.filename)
        
        if row_index < len(df):
            df.at[row_index, "Status"] = status
            df.at[row_index, "Notes"] = notes
            df.at[row_index, "Last Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
            df.to_excel(self.filename, index=False)
            self._update_status_color(row_index)
            return True
        return False
    
    def _update_status_color(self, row_index):
        """Update row color based on status"""
        wb = load_workbook(self.filename)
        ws = wb.active
        
        df = pd.read_excel(self.filename)
        status = df.at[row_index, "Status"]
        
        if status in self.status_colors:
            fill = PatternFill(start_color=self.status_colors[status], 
                              end_color=self.status_colors[status], 
                              fill_type="solid")
            
            for col in range(1, len(self.columns) + 1):
                ws.cell(row=row_index + 2, column=col).fill = fill
        
        wb.save(self.filename)
    
    def get_statistics(self):
        """Get application statistics"""
        df = pd.read_excel(self.filename)
        stats = {
            "Total": len(df),
            "Applied": len(df[df["Status"] == "Applied"]),
            "Interview": len(df[df["Status"] == "Interview"]),
            "Rejected": len(df[df["Status"] == "Rejected"]),
            "Ignored": len(df[df["Status"] == "Ignored"]),
            "Offer": len(df[df["Status"] == "Offer"])
        }
        return stats
    
    def get_applications_by_platform(self):
        """Get applications grouped by platform"""
        df = pd.read_excel(self.filename)
        return df.groupby("Platform").size().to_dict()